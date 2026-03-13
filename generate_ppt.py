from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Colors
BG = RGBColor(0x0A, 0x0A, 0x14)
SURFACE = RGBColor(0x14, 0x14, 0x22)
ACCENT = RGBColor(0x12, 0xDA, 0xA8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xA0, 0xA0, 0xB0)
DIM = RGBColor(0x66, 0x66, 0x77)
RED = RGBColor(0xEF, 0x44, 0x44)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
BLUE = RGBColor(0x3B, 0x82, 0xF6)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W = prs.slide_width
H = prs.slide_height


def set_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill_color=SURFACE, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    # Smaller corner radius
    shape.adjustments[0] = 0.05
    return shape


def add_text(slide, left, top, width, height, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox


def add_multiline(slide, left, top, width, height, lines, default_size=16, default_color=MUTED):
    """lines: list of (text, size, color, bold)"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        text, size, color, bold = line
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = 'Calibri'
        p.space_after = Pt(4)
    return txBox


def add_accent_line(slide, left, top, width):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    return shape


# =============================================
# SLIDE 1: Title
# =============================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_bg(slide)

# Accent line at top
add_accent_line(slide, Inches(0), Inches(0), W)

# Badge
add_rect(slide, Inches(5.2), Inches(2.2), Inches(3), Inches(0.4), RGBColor(0x0D, 0x2E, 0x25), ACCENT)
add_text(slide, Inches(5.2), Inches(2.2), Inches(3), Inches(0.4), "OFFLINE RETAIL FINANCING", size=11, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# Title
add_text(slide, Inches(1.5), Inches(3.0), Inches(10.3), Inches(1.2),
         "Loan MarketPlace\nfor the Physical Store", size=44, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# Subtitle
add_text(slide, Inches(2.5), Inches(4.5), Inches(8.3), Inches(0.8),
         "Eliminating the 30-minute wait for in-store financing with a self-service,\nAI-driven checkout that lives on the customer\u2019s phone.",
         size=18, color=MUTED, align=PP_ALIGN.CENTER)

# Juspay branding
add_text(slide, Inches(5.5), Inches(6.2), Inches(2.3), Inches(0.5),
         "Juspay", size=22, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)


# =============================================
# SLIDE 2: The Problem
# =============================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_text(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.3), "THE PROBLEM", size=11, color=ACCENT, bold=True)
add_text(slide, Inches(0.8), Inches(0.7), Inches(8), Inches(0.6), "The Current Offline Journey is Broken", size=32, color=WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(1.3), Inches(8), Inches(0.5),
         "In stores like Croma or Reliance Digital, financing a product is slow, manual, and frustrating.",
         size=16, color=MUTED)

# Timeline items
steps = [
    ("1. Product Selection", "Customer finds the product. Excitement is high \u2014 but the financing journey is about to begin.", None),
    ("2. Hunt for an Agent", "To explore EMI options, must locate a bank rep or store associate. Often busy with other customers.", "15\u201320 min wait"),
    ("3. Desk Dependency", "Asked to move to a dedicated desk. The entire purchase journey pauses while the loan conversation begins.", None),
    ("4. Document Hassle", "PAN Card, Aadhaar, bank details, cancelled cheque \u2014 each manually verified, photocopied, or photographed.", None),
    ("5. Manual Data Entry", "Agent manually enters information. Customer waits while details are checked and credit approval is processed.", "20\u201330 min total"),
]

left_start = Inches(0.8)
card_w = Inches(2.2)
card_h = Inches(3.5)
gap = Inches(0.2)

for i, (title, desc, tag) in enumerate(steps):
    x = left_start + i * (card_w + gap)
    y = Inches(2.2)

    card = add_rect(slide, x, y, card_w, card_h, SURFACE)

    # Red top border
    red_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, card_w, Pt(3))
    red_line.fill.solid()
    red_line.fill.fore_color.rgb = RED
    red_line.line.fill.background()

    add_text(slide, x + Inches(0.15), y + Inches(0.2), card_w - Inches(0.3), Inches(0.4), title, size=13, color=WHITE, bold=True)
    add_text(slide, x + Inches(0.15), y + Inches(0.65), card_w - Inches(0.3), Inches(2.2), desc, size=11, color=MUTED)

    if tag:
        tag_shape = add_rect(slide, x + Inches(0.15), y + card_h - Inches(0.55), Inches(1.4), Inches(0.3), RGBColor(0x2D, 0x13, 0x13), RED)
        add_text(slide, x + Inches(0.15), y + card_h - Inches(0.55), Inches(1.4), Inches(0.3), tag, size=9, color=RED, bold=True, align=PP_ALIGN.CENTER)

# Friction summary at bottom
friction_y = Inches(6.0)
add_rect(slide, Inches(0.8), friction_y, Inches(11.7), Inches(1.0), RGBColor(0x1A, 0x10, 0x10), RGBColor(0x3D, 0x15, 0x15))
add_text(slide, Inches(1.1), friction_y + Inches(0.1), Inches(3), Inches(0.3), "The Result: Excitement \u2192 Friction", size=13, color=AMBER, bold=True)

frictions = ["Agent availability dependency", "Manual documentation overhead", "Slow data entry by agents", "Multiple verification steps"]
for i, f in enumerate(frictions):
    col = i % 4
    x = Inches(1.1) + col * Inches(2.9)
    add_text(slide, x, friction_y + Inches(0.5), Inches(2.7), Inches(0.3), f"\u2022  {f}", size=11, color=MUTED)


# =============================================
# SLIDE 3: Challenge / Opportunity / Aim
# =============================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_text(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.3), "OUR APPROACH", size=11, color=ACCENT, bold=True)
add_text(slide, Inches(0.8), Inches(0.7), Inches(8), Inches(0.6), "Challenge, Opportunity & Aim", size=32, color=WHITE, bold=True)

cards_data = [
    ("The Challenge", AMBER,
     "Eliminate the friction of offline financing. Our competitors aren\u2019t just other apps \u2014 they are the physical bank representatives in-store.\n\nWe need to match that \u201chuman\u201d ease of use with a digital-first, self-service experience that feels just as supportive."),
    ("The Opportunity", BLUE,
     "Juspay has already mastered financing for e-commerce, insurance, and online education.\n\nOffline Retail is a massive untapped market. We can take the \u201cwaiting\u201d out of shopping, letting users secure financing in seconds."),
    ("Our Aim", ACCENT,
     "Create a Loan MarketPlace journey that is so frictionless it feels invisible.\n\nEmpower the customer to handle their own financing from start to finish \u2014 without needing to find a bank agent."),
]

for i, (title, color, desc) in enumerate(cards_data):
    x = Inches(0.8) + i * Inches(4.1)
    y = Inches(1.8)
    cw = Inches(3.8)
    ch = Inches(4.8)

    card = add_rect(slide, x, y, cw, ch, SURFACE)

    # Top accent line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cw, Pt(4))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()

    add_text(slide, x + Inches(0.3), y + Inches(0.4), cw - Inches(0.6), Inches(0.4), title, size=20, color=WHITE, bold=True)
    add_text(slide, x + Inches(0.3), y + Inches(1.0), cw - Inches(0.6), ch - Inches(1.4), desc, size=14, color=MUTED)


# =============================================
# SLIDE 4: Product Features
# =============================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_text(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.3), "THE PRODUCT", size=11, color=ACCENT, bold=True)
add_text(slide, Inches(0.8), Inches(0.7), Inches(10), Inches(0.6), "An AI-Driven Interface for the Physical Store", size=32, color=WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(1.3), Inches(8), Inches(0.4),
         "A smart experience on the user\u2019s smartphone that interacts with the physical retail environment.",
         size=16, color=MUTED)

features = [
    ("Product-Based QR", "Every shelf item has a unique QR code. Scanning instantly pulls up product specs and financing eligibility."),
    ("AI Recommendations", "Based on what the user scans, AI suggests similar products or better deals using semantic similarity."),
    ("Multimodal Chatbot", "Built-in assistant with typing and voice support for questions about rates, terms, and comparisons."),
    ("AI Onboarding", "Streamlined KYC, mandate setup, and automated document verification \u2014 no paper, no photocopies."),
    ("Seamless Checkout", "Initiated by user or sales agent. Guides through identity verification, mandate setup, and agreement signing."),
    ("Instant Financing", "Checks multiple lenders, lets customer pick an offer, sets up auto-pay via UPI, eNACH, or Debit Card."),
]

for i, (title, desc) in enumerate(features):
    col = i % 3
    row = i // 3
    x = Inches(0.8) + col * Inches(4.1)
    y = Inches(2.1) + row * Inches(2.5)
    cw = Inches(3.8)
    ch = Inches(2.2)

    card = add_rect(slide, x, y, cw, ch, SURFACE)

    add_text(slide, x + Inches(0.25), y + Inches(0.25), cw - Inches(0.5), Inches(0.35), title, size=16, color=WHITE, bold=True)

    # Accent underline
    uline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.25), y + Inches(0.65), Inches(0.6), Pt(2))
    uline.fill.solid()
    uline.fill.fore_color.rgb = ACCENT
    uline.line.fill.background()

    add_text(slide, x + Inches(0.25), y + Inches(0.8), cw - Inches(0.5), Inches(1.2), desc, size=12, color=MUTED)


# =============================================
# SLIDE 5: Checkout Flow
# =============================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_text(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.3), "THE JOURNEY", size=11, color=ACCENT, bold=True)
add_text(slide, Inches(0.8), Inches(0.7), Inches(8), Inches(0.6), "End-to-End Checkout Flow", size=32, color=WHITE, bold=True)

flow_steps = [
    ("1", "Checkout\nInitiation", "Initiated by user or\nsales agent via the\nmerchant platform.", []),
    ("2", "KYC\nVerification", "Multiple methods based\non lender requirements.", ["Aadhaar XML", "Aadhaar OTP", "Live Selfie"]),
    ("3", "Mandate\nSetup", "Repayment mandate for\nautomated EMI collection.", ["UPI Autopay", "eNACH"]),
    ("4", "KFS & Agreement\nSigning", "Key Fact Statement and\nloan agreement presented.", ["Clickwrap", "OTP Signing"]),
    ("5", "Disbursement", "Loan proceeds to final\nprocessing. Lender\ninitiates disbursement.", []),
]

total_w = Inches(11.7)
step_w = Inches(2.1)
arrow_w = Inches(0.3)
start_x = Inches(0.8)

for i, (num, title, desc, tags) in enumerate(flow_steps):
    x = start_x + i * (step_w + arrow_w + Inches(0.05))
    y = Inches(2.0)
    sh = Inches(4.5)

    card = add_rect(slide, x, y, step_w, sh, SURFACE)

    # Number circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.75), y + Inches(0.25), Inches(0.55), Inches(0.55))
    circle.fill.solid()
    circle.fill.fore_color.rgb = ACCENT
    circle.line.fill.background()
    tf = circle.text_frame
    tf.paragraphs[0].text = num
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BG
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.word_wrap = False

    add_text(slide, x + Inches(0.15), y + Inches(1.0), step_w - Inches(0.3), Inches(0.7), title, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.15), y + Inches(1.8), step_w - Inches(0.3), Inches(1.2), desc, size=11, color=MUTED, align=PP_ALIGN.CENTER)

    # Tags
    tag_y = y + Inches(3.2)
    for j, tag in enumerate(tags):
        tag_shape = add_rect(slide, x + Inches(0.2), tag_y + j * Inches(0.35), step_w - Inches(0.4), Inches(0.28),
                            RGBColor(0x0D, 0x2E, 0x25), None)
        add_text(slide, x + Inches(0.2), tag_y + j * Inches(0.35), step_w - Inches(0.4), Inches(0.28),
                tag, size=9, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

    # Arrow between steps
    if i < len(flow_steps) - 1:
        ax = x + step_w + Inches(0.05)
        ay = y + sh / 2 - Inches(0.1)
        add_text(slide, ax, ay, arrow_w, Inches(0.3), "\u2192", size=20, color=DIM, align=PP_ALIGN.CENTER)


# =============================================
# SLIDE 6: Target Audience
# =============================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_text(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.3), "WHO IT\u2019S FOR", size=11, color=ACCENT, bold=True)
add_text(slide, Inches(0.8), Inches(0.7), Inches(8), Inches(0.6), "Target Audience", size=32, color=WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(1.3), Inches(8), Inches(0.4),
         "Designed to be accessible across the entire technical spectrum.",
         size=16, color=MUTED)

audiences = [
    ("The Tech-Savvy Buyer", "POWER USER",
     "Enjoys the speed, the data-rich UI, and the ability to bypass human interaction entirely.\n\n\u2022  Compare offers side-by-side\n\u2022  View detailed product specs\n\u2022  Complete checkout autonomously\n\u2022  Self-service from start to finish"),
    ("The Non-Tech Buyer", "GUIDED EXPERIENCE",
     "Benefits from voice chat and AI guidance that acts like a \u201cdigital store assistant.\u201d\n\n\u2022  No app downloads needed\n\u2022  Voice-first interaction\n\u2022  Step-by-step guidance\n\u2022  Familiar and simple experience"),
]

for i, (title, persona, desc) in enumerate(audiences):
    x = Inches(0.8) + i * Inches(6.1)
    y = Inches(2.2)
    cw = Inches(5.8)
    ch = Inches(4.5)

    card = add_rect(slide, x, y, cw, ch, SURFACE)

    # Top accent
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cw, Pt(4))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    add_text(slide, x + Inches(0.4), y + Inches(0.4), cw - Inches(0.8), Inches(0.45), title, size=22, color=WHITE, bold=True)

    # Persona badge
    badge = add_rect(slide, x + Inches(0.4), y + Inches(0.95), Inches(1.8), Inches(0.3), RGBColor(0x0D, 0x2E, 0x25))
    add_text(slide, x + Inches(0.4), y + Inches(0.95), Inches(1.8), Inches(0.3), persona, size=9, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

    add_text(slide, x + Inches(0.4), y + Inches(1.5), cw - Inches(0.8), Inches(2.8), desc, size=13, color=MUTED)


# =============================================
# SLIDE 7: Revenue Model
# =============================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_text(slide, Inches(0.8), Inches(0.4), Inches(4), Inches(0.3), "BUSINESS MODEL", size=11, color=ACCENT, bold=True)
add_text(slide, Inches(0.8), Inches(0.7), Inches(8), Inches(0.6), "How Juspay Generates Revenue", size=32, color=WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(1.3), Inches(8), Inches(0.4),
         "Revenue built on transaction success and the value we provide to lenders and retailers.",
         size=16, color=MUTED)

revenues = [
    ("Transaction Fees", ACCENT,
     "A small fee for every successful loan processed through our platform.",
     "Per-transaction revenue that scales with volume."),
    ("Lender Integration", BLUE,
     "Banks and NBFCs pay to integrate their financing products into our high-speed journey.",
     "Recurring integration and maintenance fees."),
    ("Technology Licensing", AMBER,
     "Retailers pay for the platform\u2019s ability to increase conversion rates.",
     "Turns window shoppers into buyers through instant credit."),
]

for i, (title, color, desc, note) in enumerate(revenues):
    x = Inches(0.8) + i * Inches(4.1)
    y = Inches(2.2)
    cw = Inches(3.8)
    ch = Inches(4.3)

    card = add_rect(slide, x, y, cw, ch, SURFACE)

    # Color top line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cw, Pt(4))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()

    add_text(slide, x + Inches(0.3), y + Inches(0.5), cw - Inches(0.6), Inches(0.4), title, size=20, color=WHITE, bold=True)
    add_text(slide, x + Inches(0.3), y + Inches(1.2), cw - Inches(0.6), Inches(1.5), desc, size=14, color=MUTED)

    # Note at bottom
    note_bg = add_rect(slide, x + Inches(0.2), y + ch - Inches(1.0), cw - Inches(0.4), Inches(0.6), RGBColor(0x0D, 0x2E, 0x25))
    add_text(slide, x + Inches(0.35), y + ch - Inches(0.95), cw - Inches(0.7), Inches(0.5), note, size=11, color=ACCENT)


# =============================================
# SLIDE 8: Closing
# =============================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_accent_line(slide, Inches(0), Inches(0), W)

add_text(slide, Inches(1.5), Inches(2.5), Inches(10.3), Inches(1),
         "Making Offline Financing\nFeel Invisible", size=44, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

add_text(slide, Inches(2.5), Inches(4.2), Inches(8.3), Inches(0.6),
         "From 30-minute waits to instant, self-service checkout.\nPowered by QR, AI, and seamless lending infrastructure.",
         size=18, color=MUTED, align=PP_ALIGN.CENTER)

# Juspay
add_text(slide, Inches(5.5), Inches(5.5), Inches(2.3), Inches(0.5),
         "Juspay", size=24, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(3.5), Inches(6.0), Inches(6.3), Inches(0.4),
         "Powering frictionless financing for the next billion transactions.",
         size=13, color=DIM, align=PP_ALIGN.CENTER)


# Save
output = "/Users/prakhar.prakash/Desktop/Clones/LMP/LMP_Product_Document.pptx"
prs.save(output)
print(f"Saved: {output}")
